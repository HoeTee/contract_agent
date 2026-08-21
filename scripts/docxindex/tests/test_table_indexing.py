from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

PROJECT_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = PROJECT_DIR.parents[1]
sys.path.insert(0, str(PROJECT_DIR))
sys.path.insert(0, str(PROJECT_ROOT))

import cli
from docxindex import build_document_index
from docxindex.retrieval import build_content_context


class TableIndexingTest(unittest.TestCase):
    def _write_contract(self, path: Path) -> None:
        doc = Document()
        doc.add_paragraph("测试采购合同")
        doc.add_paragraph("第一条 合同金额")
        doc.add_paragraph("合同总价为人民币300元。")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "金额"
        table.cell(1, 0).text = "设备"
        table.cell(1, 1).text = "100"
        table.cell(2, 0).text = "服务"
        table.cell(2, 1).text = "200"
        doc.add_paragraph("第二条 其他约定")
        doc.add_paragraph("双方共同遵守。")
        doc.save(path)

    def test_table_is_independent_node_and_parent_does_not_duplicate_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx = Path(directory) / "contract.docx"
            self._write_contract(docx)

            index = build_document_index(docx).to_json_dict()
            table_nodes = [node for node in index["nodes"] if node["node_type"] == "table"]

            self.assertEqual(len(table_nodes), 1)
            table_node = table_nodes[0]
            self.assertIn("| 项目 [S0001] | 金额 [S0002] |", table_node["text"])
            self.assertEqual(table_node["start_anchor"], table_node["end_anchor"])
            self.assertEqual(table_node["mapping_ref"], f"table_store.json#{table_node['table_id']}")
            parent = next(node for node in index["nodes"] if node["node_id"] == table_node["parent_id"])
            self.assertNotIn("| 项目", parent["text"])
            self.assertIn(table_node["table_id"], index["table_map"])
            self.assertEqual(index["table_map"][table_node["table_id"]]["rows"], 3)

    def test_cli_persists_table_store(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docx = root / "contract.docx"
            output = root / "outputs"
            self._write_contract(docx)

            result = cli.write_document_outputs(
                docx,
                output,
                [],
                "llm",
                None,
                False,
                False,
                None,
                False,
                None,
            )
            document_dir = Path(result["output_dir"])
            manifest = json.loads((document_dir / "document_index.json").read_text(encoding="utf-8"))
            table_store = json.loads((document_dir / "table_store.json").read_text(encoding="utf-8"))

            self.assertEqual(manifest["stats"]["table_count"], 1)
            self.assertEqual(len(table_store), 1)
            mapping = next(iter(table_store.values()))
            self.assertIn("markdown", mapping)
            self.assertIn("sources", mapping)

    def test_oversized_table_parent_expands_to_row_chunks(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx = Path(directory) / "contract.docx"
            self._write_contract(docx)
            index = build_document_index(
                docx,
                table_split_threshold_tokens=20,
                table_chunk_target_tokens=15,
            ).to_json_dict()
            table_node = next(node for node in index["nodes"] if node["node_type"] == "table")

            self.assertGreater(len(table_node["children"]), 0)
            context = build_content_context(
                index,
                [{"node_id": table_node["node_id"], "source": "structure"}],
                input_tokens=1000,
            )

            returned = context["content_context"]
            self.assertEqual([item["node_id"] for item in returned], table_node["children"])
            self.assertTrue(all(item["text"].startswith("<!-- table_id:") for item in returned))
            self.assertTrue(all(item["table_id"] == table_node["table_id"] for item in returned))
            self.assertTrue(all(item["mapping_ref"] == table_node["mapping_ref"] for item in returned))

    def test_table_node_is_atomic_above_normal_content_budget(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx = Path(directory) / "contract.docx"
            self._write_contract(docx)
            index = build_document_index(docx).to_json_dict()
            table_node = next(node for node in index["nodes"] if node["node_type"] == "table")

            context = build_content_context(
                index,
                [{"node_id": table_node["node_id"], "source": "structure"}],
                input_tokens=10,
            )

            returned = context["content_context"]
            self.assertEqual(len(returned), 1)
            self.assertEqual(returned[0]["node_id"], table_node["node_id"])
            self.assertEqual(returned[0]["text"], table_node["text"])
            self.assertFalse(returned[0]["truncated"])


if __name__ == "__main__":
    unittest.main()

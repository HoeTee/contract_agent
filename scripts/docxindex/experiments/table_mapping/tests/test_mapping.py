from __future__ import annotations

import tempfile
import unittest
import sys
from pathlib import Path

DOCXINDEX_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(DOCXINDEX_ROOT))

from docx import Document

from docxindex.table_mapping import (
    cell_source_refs,
    render_table_markdown,
    resolve_agent_evidence,
)


class TableMarkdownMappingTest(unittest.TestCase):
    def test_merged_cell_reuses_source_and_unique_quote_resolves(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "merged.docx"
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            table.cell(0, 0).text = "类别"
            table.cell(0, 1).text = "金额"
            merged = table.cell(1, 0).merge(table.cell(2, 0))
            merged.text = "设备类"
            table.cell(1, 1).text = "925129"
            table.cell(2, 1).text = "842420"
            doc.save(path)

            loaded = Document(path)
            mapping = render_table_markdown(loaded.tables[0], "body/tbl1")
            self.assertEqual(cell_source_refs(mapping, 2, 1), cell_source_refs(mapping, 3, 1))
            self.assertIn("同上", mapping.markdown)

            resolved = resolve_agent_evidence(mapping, "925129")
            self.assertEqual(resolved.status, "matched")
            self.assertEqual(resolved.matched_original_text, "925129")

    def test_horizontal_and_vertical_merges_restore_logical_grid(self) -> None:
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        horizontal = table.cell(0, 0).merge(table.cell(0, 1))
        horizontal.text = "合并表头"
        table.cell(0, 2).text = "金额"
        vertical = table.cell(1, 0).merge(table.cell(2, 0))
        vertical.text = "设备类"
        table.cell(1, 1).text = "服务器"
        table.cell(1, 2).text = "100"
        table.cell(2, 1).text = "交换机"
        table.cell(2, 2).text = "200"

        mapping = render_table_markdown(table, "body/tbl1")

        self.assertEqual(mapping.rows, 3)
        self.assertEqual(mapping.columns, 3)
        self.assertIn("同左", mapping.markdown)
        self.assertIn("同上", mapping.markdown)
        self.assertEqual(cell_source_refs(mapping, 1, 1), cell_source_refs(mapping, 1, 2))
        self.assertEqual(cell_source_refs(mapping, 2, 1), cell_source_refs(mapping, 3, 1))
        merged_cell = next(cell for cell in mapping.cells if cell.row == 2 and cell.column == 1)
        self.assertEqual((merged_cell.row_start, merged_cell.row_end), (2, 3))

    def test_repeated_quote_is_ambiguous_but_source_ref_is_exact(self) -> None:
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "数量"
        table.cell(0, 1).text = "数量"
        table.cell(1, 0).text = "2"
        table.cell(1, 1).text = "2"
        mapping = render_table_markdown(table, "body/tbl1")

        self.assertEqual(resolve_agent_evidence(mapping, "2").status, "ambiguous")
        source_ref = cell_source_refs(mapping, 2, 1)[0]
        exact = resolve_agent_evidence(mapping, "2", source_ref)
        self.assertEqual(exact.status, "matched")

    def test_source_mismatch_is_rejected(self) -> None:
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "925129"
        mapping = render_table_markdown(table, "body/tbl1")
        source_ref = cell_source_refs(mapping, 1, 1)[0]
        self.assertEqual(resolve_agent_evidence(mapping, "842420", source_ref).status, "source_mismatch")

    def test_normalized_quote_returns_original_text(self) -> None:
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "190968"
        mapping = render_table_markdown(table, "body/tbl1")
        source_ref = cell_source_refs(mapping, 1, 1)[0]
        resolved = resolve_agent_evidence(mapping, "190 968", source_ref)
        self.assertEqual(resolved.status, "matched")
        self.assertEqual(resolved.matched_original_text, "190968")


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from tools.document.table_markdown_map import (
    cell_source_refs,
    render_table_markdown,
    resolve_agent_evidence,
)


class TableMarkdownMappingTest(unittest.TestCase):
    def test_merged_cell_reuses_source_and_unique_quote_resolves(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "merged.docx"
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            table.cell(0, 0).text = "类别"
            table.cell(0, 1).text = "金额"
            merged = table.cell(1, 0).merge(table.cell(2, 0))
            merged.text = "设备类"
            table.cell(1, 1).text = "925129"
            table.cell(2, 1).text = "842420"
            doc.save(path)

            loaded = Document(path)
            mapping = render_table_markdown(loaded.tables[0], "body/tbl1")
            self.assertEqual(cell_source_refs(mapping, 2, 1), cell_source_refs(mapping, 3, 1))
            self.assertIn("同上", mapping.markdown)

            resolved = resolve_agent_evidence(mapping, "925129")
            self.assertEqual(resolved.status, "matched")
            self.assertEqual(resolved.matched_original_text, "925129")

    def test_repeated_quote_is_ambiguous_but_source_ref_is_exact(self) -> None:
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "数量"
        table.cell(0, 1).text = "数量"
        table.cell(1, 0).text = "2"
        table.cell(1, 1).text = "2"
        mapping = render_table_markdown(table, "body/tbl1")

        self.assertEqual(resolve_agent_evidence(mapping, "2").status, "ambiguous")
        source_ref = cell_source_refs(mapping, 2, 1)[0]
        exact = resolve_agent_evidence(mapping, "2", source_ref)
        self.assertEqual(exact.status, "matched")

    def test_source_mismatch_is_rejected(self) -> None:
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "925129"
        mapping = render_table_markdown(table, "body/tbl1")
        source_ref = cell_source_refs(mapping, 1, 1)[0]
        self.assertEqual(resolve_agent_evidence(mapping, "842420", source_ref).status, "source_mismatch")

    def test_normalized_quote_returns_original_text(self) -> None:
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "190968"
        mapping = render_table_markdown(table, "body/tbl1")
        source_ref = cell_source_refs(mapping, 1, 1)[0]
        resolved = resolve_agent_evidence(mapping, "190 968", source_ref)
        self.assertEqual(resolved.status, "matched")
        self.assertEqual(resolved.matched_original_text, "190968")


if __name__ == "__main__":
    unittest.main()

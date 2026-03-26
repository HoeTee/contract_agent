import os
import re
import copy
import shutil
import tempfile
import zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pypdf import PdfReader
from datetime import datetime
from lxml import etree
from tools.clean_docx import DocxCleaner


class FileParser:
    """Parses DOCX, PDF, and TXT files into text chunks."""
    
    @staticmethod
    def parse_file(file_path: str) -> str:
        """Determines file type and extracts text."""
        ext = os.path.splitext(file_path)[1].lower()
        
        if not os.path.exists(file_path):
            return f"Error: File not found at {file_path}"
            
        try:
            if ext == '.docx':
                clean_path = FileParser._clean_docx(file_path)
                try:
                    return FileParser._parse_docx(clean_path)
                finally:
                    if clean_path != file_path and os.path.exists(clean_path):
                        os.unlink(clean_path)
            elif ext == '.pdf':
                return FileParser._parse_pdf(file_path)
            elif ext == '.txt':
                return FileParser._parse_txt(file_path)
            else:
                return f"Error: Unsupported file format {ext}"
        except Exception as e:
            return f"Error parsing file: {str(e)}"

    @staticmethod
    def _clean_docx(path: str) -> str:
        cleaner = DocxCleaner(input_path=path)
        return cleaner.clean() 

    @staticmethod
    def _clean_docx_for_parsing(path: str) -> str:
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        tmp.close()

        def _unwrap(root, tag: str) -> None:
            for el in root.xpath(f".//w:{tag}", namespaces=ns):
                parent = el.getparent()
                if parent is None:
                    continue
                idx = parent.index(el)
                for child in list(el):
                    parent.insert(idx, child)
                    idx += 1
                parent.remove(el)

        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp.name, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    root = etree.fromstring(data)
                    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference", "moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd"):
                        for el in root.xpath(f".//w:{tag}", namespaces=ns):
                            parent = el.getparent()
                            if parent is not None:
                                parent.remove(el)
                    _unwrap(root, "ins")
                    _unwrap(root, "moveTo")
                    for tag in ("del", "moveFrom"):
                        for el in root.xpath(f".//w:{tag}", namespaces=ns):
                            parent = el.getparent()
                            if parent is not None:
                                parent.remove(el)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                zout.writestr(info, data)

        return tmp.name

    @staticmethod
    def _parse_docx(path: str) -> str:
        doc = Document(path)
        preamble = []   # Content before the first heading
        full_text = []
        found_first_heading = False
        first_heading_level = 1  # Default if no headings found
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            # Convert DOCX heading styles to markdown
            style = para.style.name if para.style else ""
            if style.startswith("Heading"):
                try:
                    level = int(style.split()[-1])
                    level = min(level, 6)
                except (ValueError, IndexError):
                    level = 1
                if not found_first_heading:
                    first_heading_level = level
                    if preamble:
                        # Use same level as first heading so preamble is a sibling
                        full_text.append(f"{'#' * level} 合同首部信息")
                        full_text.extend(preamble)
                    found_first_heading = True
                full_text.append(f"{'#' * level} {text}")
            else:
                if found_first_heading:
                    full_text.append(text)
                else:
                    preamble.append(text)
        # If the document has no headings at all, emit preamble under a single heading
        if not found_first_heading and preamble:
            full_text.append("# 合同首部信息")
            full_text.extend(preamble)
        return "\n\n".join(full_text)

    @staticmethod
    def _parse_pdf(path: str) -> str:
        reader = PdfReader(path)
        full_text = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                full_text.append(text)
        return "\n".join(full_text)

    @staticmethod
    def _parse_txt(path: str) -> str:
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()


class ReportGenerator:
    """Generates MD, DOCX (with comments on original contract), and PDF reports.
    
    Pipeline:
    - MD: saved as-is
    - DOCX: original contract + comments (annotations) from review results
    - PDF: Markdown → HTML → PDF via xhtml2pdf with optimized CSS
    """

    # CSS for PDF rendering with optimized Chinese support and table layout
    _PDF_CSS = """\
@page {
    size: A4;
    margin: 20mm 15mm 20mm 15mm;
}
body {
    font-family: "SimSun", "SimHei", "Microsoft YaHei", serif;
    font-size: 11pt;
    line-height: 1.6;
    color: #222;
}
h1 {
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
    font-size: 18pt;
    border-bottom: 2px solid #333;
    padding-bottom: 4px;
    margin-top: 20px;
    page-break-after: avoid;
}
h2 {
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
    font-size: 15pt;
    border-bottom: 1px solid #aaa;
    padding-bottom: 3px;
    margin-top: 16px;
    page-break-after: avoid;
}
h3 {
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
    font-size: 13pt;
    margin-top: 14px;
    page-break-after: avoid;
}
h4 {
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
    font-size: 12pt;
    margin-top: 12px;
    page-break-after: avoid;
}
table {
    border-collapse: collapse;
    width: 100%;
    margin: 10px 0;
    table-layout: fixed;
    page-break-inside: avoid;
}
th, td {
    border: 1px solid #999;
    padding: 4px 6px;
    text-align: left;
    font-size: 9pt;
    word-wrap: break-word;
    overflow-wrap: break-word;
    word-break: break-all;
}
th {
    background-color: #f0f0f0;
    font-weight: bold;
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
}
blockquote {
    border-left: 3px solid #ccc;
    margin: 8px 0;
    padding: 6px 12px;
    color: #444;
    background-color: #f9f9f9;
    font-size: 10pt;
    page-break-inside: avoid;
}
code {
    background-color: #f4f4f4;
    padding: 1px 3px;
    border-radius: 2px;
    font-size: 9pt;
}
pre {
    background-color: #f4f4f4;
    padding: 8px;
    border-radius: 3px;
    overflow-x: auto;
    page-break-inside: avoid;
}
pre code {
    padding: 0;
    background: none;
}
hr {
    border: none;
    border-top: 1px solid #ccc;
    margin: 16px 0;
}
ul, ol { padding-left: 20px; }
li { margin-bottom: 3px; font-size: 11pt; }
strong { font-weight: bold; }
em { font-style: italic; }
p { margin: 4px 0; }
"""

    # ==================== DOCX Comment Helpers ====================

    @staticmethod
    def _normalize_text(text: str) -> str:
        """Normalize text for fuzzy matching: strip whitespace, punctuation variations."""
        if not text:
            return ""
        # Remove markdown quote markers
        text = re.sub(r'^>\s*', '', text, flags=re.MULTILINE)
        # Collapse all whitespace (including \u3000 fullwidth space)
        text = re.sub(r'[\s\u3000]+', '', text)
        # Normalize some common punctuation variants
        text = text.replace('（', '(').replace('）', ')').replace('：', ':')
        text = text.replace('，', ',').replace('。', '.').replace('；', ';')
        return text

    @staticmethod
    def _find_matching_paragraphs(doc: Document, reference_text: str) -> list:
        """Find paragraphs in doc that contain fragments of the reference text.
        
        Returns a list of paragraph indices that match.
        """
        if not reference_text or not reference_text.strip():
            return []

        # Clean reference text (remove markdown blockquote markers etc.)
        clean_ref = re.sub(r'^>\s*', '', reference_text, flags=re.MULTILINE)
        # Split reference into meaningful lines (skip empty)
        ref_lines = [l.strip() for l in clean_ref.split('\n') if l.strip()]
        
        if not ref_lines:
            return []

        matched_indices = []
        for i, para in enumerate(doc.paragraphs):
            para_text = para.text.strip()
            if not para_text or len(para_text) < 4:
                continue
            
            norm_para = ReportGenerator._normalize_text(para_text)
            if not norm_para:
                continue
            
            # Check if any reference line matches this paragraph
            for ref_line in ref_lines:
                norm_ref = ReportGenerator._normalize_text(ref_line)
                if not norm_ref or len(norm_ref) < 4:
                    continue
                # Check substring match (either direction for flexibility)
                if norm_ref in norm_para or norm_para in norm_ref:
                    if i not in matched_indices:
                        matched_indices.append(i)
                    break
                # Also check significant overlap (first 20 chars)
                check_len = min(20, len(norm_ref), len(norm_para))
                if check_len >= 8 and norm_ref[:check_len] == norm_para[:check_len]:
                    if i not in matched_indices:
                        matched_indices.append(i)
                    break

        return matched_indices

    @staticmethod
    def _add_comments_to_doc(doc: Document, comments_data: list) -> Document:
        """Add Word comments to a document using OPC XML manipulation.
        
        Args:
            doc: The python-docx Document object.
            comments_data: List of dicts with keys:
                - 'paragraph_indices': list of int (paragraph indices to annotate)
                - 'comment_text': str (the comment content)
                - 'author': str
        """
        # Get or create the comments part
        comments_part = None
        for rel in doc.part.rels.values():
            if "comments" in rel.reltype:
                comments_part = rel.target_part
                break

        if comments_part is None:
            # Create new comments XML part
            from docx.opc.part import Part
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            
            comments_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
                ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '</w:comments>'
            )
            comments_part = Part(
                partname="/word/comments.xml",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
                blob=comments_xml.encode('utf-8'),
                package=doc.part.package,
            )
            doc.part.relate_to(comments_part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
        
        # Parse the comments XML
        comments_element = etree.fromstring(comments_part.blob)
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        
        comment_id = 0
        now_str = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")

        for cdata in comments_data:
            para_indices = cdata.get('paragraph_indices', [])
            comment_text = cdata.get('comment_text', '')
            author = cdata.get('author', 'AI审查助手')

            if not para_indices or not comment_text:
                continue

            # Create <w:comment> element
            comment_el = etree.SubElement(comments_element, qn('w:comment'))
            comment_el.set(qn('w:id'), str(comment_id))
            comment_el.set(qn('w:author'), author)
            comment_el.set(qn('w:date'), now_str + "Z")

            # Split comment text into paragraphs for readability
            comment_lines = comment_text.split('\n')
            for line in comment_lines:
                line = line.strip()
                if not line:
                    continue
                p_el = etree.SubElement(comment_el, qn('w:p'))
                r_el = etree.SubElement(p_el, qn('w:r'))
                t_el = etree.SubElement(r_el, qn('w:t'))
                t_el.set(qn('xml:space'), 'preserve')
                t_el.text = line

            # Add commentRangeStart and commentRangeEnd to the first matched paragraph
            target_idx = para_indices[0]
            if target_idx < len(doc.paragraphs):
                para_element = doc.paragraphs[target_idx]._element

                # Insert commentRangeStart at beginning
                range_start = OxmlElement('w:commentRangeStart')
                range_start.set(qn('w:id'), str(comment_id))
                para_element.insert(0, range_start)

                # Append commentRangeEnd and commentReference at end
                range_end = OxmlElement('w:commentRangeEnd')
                range_end.set(qn('w:id'), str(comment_id))
                para_element.append(range_end)

                # Add comment reference run
                ref_run = OxmlElement('w:r')
                ref_rpr = OxmlElement('w:rPr')
                ref_style = OxmlElement('w:rStyle')
                ref_style.set(qn('w:val'), 'CommentReference')
                ref_rpr.append(ref_style)
                ref_run.append(ref_rpr)
                ref_mark = OxmlElement('w:commentReference')
                ref_mark.set(qn('w:id'), str(comment_id))
                ref_run.append(ref_mark)
                para_element.append(ref_run)

            comment_id += 1

        # Write updated comments XML back
        comments_part._blob = etree.tostring(comments_element, xml_declaration=True, encoding='UTF-8', standalone=True)
        
        return doc

    @staticmethod
    def _parse_review_output(review_output: str) -> dict:
        """Parse a sub-agent's review_output markdown into structured fields.
        
        Extracts: 风险等级, 审查结论, 原文引用, 问题分析, 修改建议
        """
        fields = {}
        
        # Pattern: "- **字段名**：内容" or "- **字段名**：\n> 引用内容"
        # Also handle multi-line content
        current_field = None
        current_content = []
        
        for line in review_output.split('\n'):
            # Match field headers like "- **风险等级**：高"
            m = re.match(r'^-\s*\*\*(.+?)\*\*\s*[：:]\s*(.*)', line)
            if m:
                # Save previous field
                if current_field:
                    fields[current_field] = '\n'.join(current_content).strip()
                current_field = m.group(1).strip()
                first_val = m.group(2).strip()
                current_content = [first_val] if first_val else []
            elif current_field:
                current_content.append(line)
        
        # Save last field
        if current_field:
            fields[current_field] = '\n'.join(current_content).strip()
        
        return fields

    @staticmethod
    def _build_comment_text(parsed: dict) -> str:
        """Build comment text from parsed review fields."""
        parts = []
        
        risk = parsed.get('风险等级', '').strip()
        if risk:
            parts.append(f"【风险等级：{risk}】")
        
        conclusion = parsed.get('审查结论', '').strip()
        if conclusion:
            parts.append(f"审查结论：{conclusion}")
        
        analysis = parsed.get('问题分析', '').strip()
        if analysis:
            parts.append(f"问题分析：{analysis}")
        
        suggestion = parsed.get('修改建议', '').strip()
        if suggestion:
            # Clean markdown blockquote markers from suggestion
            suggestion = re.sub(r'^>\s*', '', suggestion, flags=re.MULTILINE)
            parts.append(f"修改建议：{suggestion}")
        
        return '\n'.join(parts)

    @staticmethod
    def _generate_docx_with_comments(
        contract_path: str,
        results: list,
        output_path: str,
    ) -> str:
        """Generate a DOCX by copying the original contract and adding review comments.
        
        Args:
            contract_path: Path to original contract DOCX file.
            results: List of result dicts from orchestrator (with review_output).
            output_path: Where to save the annotated DOCX.
            
        Returns:
            The output path if successful, None if failed.
        """
        if not os.path.exists(contract_path):
            print(f"  WARNING: Contract file not found: {contract_path}")
            return None
        
        # Ensure contract is a DOCX file
        if not contract_path.lower().endswith('.docx'):
            print(f"  WARNING: Contract is not a DOCX file: {contract_path}")
            return None

        try:
            # Copy contract as base
            shutil.copy2(contract_path, output_path)
            doc = Document(output_path)
            
            comments_data = []
            unmatched_comments = []
            
            for result in results:
                review_output = result.get('review_output', '').strip()
                if not review_output:
                    continue  # COMPLIANT or ERROR, skip
                
                criterion = result.get('criterion', '')
                cid = result.get('criterion_id', '')
                
                # The review_output may contain multiple check points (#### sections)
                # Split by check point headers
                sections = re.split(r'(?=^####\s)', review_output, flags=re.MULTILINE)
                
                for section in sections:
                    section = section.strip()
                    if not section:
                        continue
                    
                    parsed = ReportGenerator._parse_review_output(section)
                    reference = parsed.get('原文引用', '')
                    comment_text = ReportGenerator._build_comment_text(parsed)
                    
                    if not comment_text:
                        continue
                    
                    # Extract check point title if present
                    title_match = re.match(r'^####\s*(.+)', section)
                    check_title = title_match.group(1).strip() if title_match else criterion
                    comment_text = f"[{check_title}]\n{comment_text}"
                    
                    # Try to find matching paragraphs
                    matched = ReportGenerator._find_matching_paragraphs(doc, reference)
                    
                    if matched:
                        comments_data.append({
                            'paragraph_indices': matched,
                            'comment_text': comment_text,
                            'author': 'AI审查助手',
                        })
                    else:
                        # Collect unmatched for appending at end
                        unmatched_comments.append({
                            'comment_text': comment_text,
                            'author': 'AI审查助手',
                        })
            
            # Add unmatched comments at document end
            if unmatched_comments:
                # Add a separator paragraph
                sep_para = doc.add_paragraph()
                sep_para.add_run("─" * 40)
                
                note_para = doc.add_paragraph()
                note_para.add_run("【以下为未能匹配到原文位置的审查批注】").bold = True
                
                for uc in unmatched_comments:
                    end_para = doc.add_paragraph()
                    end_para.add_run(uc['comment_text'][:50] + "...")
                    # Add comment on this trailing paragraph
                    comments_data.append({
                        'paragraph_indices': [len(doc.paragraphs) - 1],
                        'comment_text': uc['comment_text'],
                        'author': uc['author'],
                    })
            
            # Apply all comments to the document
            if comments_data:
                ReportGenerator._add_comments_to_doc(doc, comments_data)
            
            doc.save(output_path)
            matched_count = len(comments_data) - len(unmatched_comments)
            print(f"  DOCX: {len(comments_data)} comments ({matched_count} matched, {len(unmatched_comments)} at end)")
            return output_path
            
        except Exception as e:
            print(f"  WARNING: DOCX comment generation failed: {e}")
            import traceback
            traceback.print_exc()
            return None

    # ==================== PDF Generation ====================

    @staticmethod
    def _markdown_to_html(md_content: str) -> str:
        """Convert markdown text to an HTML body fragment."""
        import markdown
        extensions = [
            'tables',
            'fenced_code',
            'codehilite',
            'toc',
            'nl2br',
            'sane_lists',
        ]
        extension_configs = {
            'codehilite': {'css_class': 'highlight', 'guess_lang': False},
        }
        html_body = markdown.markdown(
            md_content,
            extensions=extensions,
            extension_configs=extension_configs,
        )
        return html_body

    @staticmethod
    def _wrap_html_for_pdf(html_body: str) -> str:
        """Wrap an HTML body fragment in a full HTML document with PDF-optimized CSS."""
        return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<style>
{ReportGenerator._PDF_CSS}
</style>
</head>
<body>
{html_body}
</body>
</html>"""

    @staticmethod
    def _register_chinese_fonts():
        """Register system Chinese TTF fonts for xhtml2pdf/reportlab."""
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        
        # Common Windows font paths
        font_dir = r"C:\Windows\Fonts"
        fonts_to_register = [
            ("SimSun", "simsun.ttc"),
            ("SimHei", "simhei.ttf"),
            ("Microsoft-YaHei", "msyh.ttc"),
        ]
        
        registered = []
        for font_name, font_file in fonts_to_register:
            font_path = os.path.join(font_dir, font_file)
            if os.path.exists(font_path):
                try:
                    pdfmetrics.registerFont(TTFont(font_name, font_path))
                    registered.append(font_name)
                except Exception as e:
                    print(f"  WARNING: Could not register font {font_name}: {e}")
        
        if registered:
            print(f"  PDF fonts registered: {', '.join(registered)}")
        else:
            # Fallback: register CID font
            print("  WARNING: No TTF fonts found, falling back to CID font")
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))

    # ==================== Main Report Generator ====================

    @staticmethod
    def generate_report(
        content: str,
        contract_name: str,
        elapsed_seconds: float = None,
        contract_path: str = None,
        results: list = None,
    ) -> dict:
        """
        Generate MD, DOCX, and PDF reports from markdown content.

        Each format is saved in its own folder under docs/:
          - docs/reports_md/
          - docs/reports_docx/
          - docs/reports_pdf/

        Args:
            content: The report body in markdown format.
            contract_name: Name of the contract (used in filenames).
            elapsed_seconds: Total workflow elapsed time in seconds (optional).
            contract_path: Path to original contract DOCX (for annotated DOCX output).
            results: List of structured review results from orchestrator.

        Returns:
            dict with keys 'md', 'docx', 'pdf' mapping to absolute file paths.
        """
        project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

        # Separate output directories for each format
        md_dir = os.path.join(project_root, "docs", "reports_md")
        docx_dir = os.path.join(project_root, "docs", "reports_docx")
        pdf_dir = os.path.join(project_root, "docs", "reports_pdf")
        for d in (md_dir, docx_dir, pdf_dir):
            os.makedirs(d, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = f"审查报告_{contract_name}_{timestamp}"

        # Append elapsed time info to the content if provided
        if elapsed_seconds is not None:
            mins, secs = divmod(int(elapsed_seconds), 60)
            time_info = (
                f"\n\n---\n\n"
                f"## 附：报告生成信息\n\n"
                f"- **总耗时**: {mins}m {secs}s\n"
                f"- **生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
            )
            content += time_info

        # 1. Save MD
        md_path = os.path.join(md_dir, f"{base_name}.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(content)

        # 2. Generate DOCX with comments on original contract
        docx_path = os.path.join(docx_dir, f"{base_name}.docx")
        if contract_path and results:
            docx_path = ReportGenerator._generate_docx_with_comments(
                contract_path, results, docx_path
            )
        else:
            # Fallback: plain DOCX from HTML if no contract path
            try:
                from htmldocx import HtmlToDocx
                html_body = ReportGenerator._markdown_to_html(content)
                full_html = f'<html><body>{html_body}</body></html>'
                doc = Document()
                parser = HtmlToDocx()
                parser.add_html_to_document(full_html, doc)
                doc.save(docx_path)
            except Exception as e:
                print(f"  WARNING: DOCX generation failed: {e}")
                docx_path = None

        # 3. Generate PDF from HTML with optimized CSS
        html_body = ReportGenerator._markdown_to_html(content)
        pdf_html = ReportGenerator._wrap_html_for_pdf(html_body)
        pdf_path = os.path.join(pdf_dir, f"{base_name}.pdf")
        try:
            from xhtml2pdf import pisa
            
            # Register Chinese TTF fonts
            ReportGenerator._register_chinese_fonts()
            
            with open(pdf_path, "wb") as pdf_file:
                pisa_status = pisa.CreatePDF(pdf_html, dest=pdf_file)
                if pisa_status.err:
                    print(f"  WARNING: PDF generation had errors (code {pisa_status.err})")
                    pdf_path = None
        except Exception as e:
            print(f"  WARNING: PDF generation failed: {e}")
            pdf_path = None

        return {
            "md": os.path.abspath(md_path),
            "docx": os.path.abspath(docx_path) if docx_path else None,
            "pdf": os.path.abspath(pdf_path) if pdf_path else None,
        }

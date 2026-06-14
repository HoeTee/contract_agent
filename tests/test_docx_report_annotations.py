import os
import json
import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

MODULE_PATH = Path(__file__).resolve().parents[1] / "tools" / "document" / "reporting" / "docx_report.py"
SPEC = importlib.util.spec_from_file_location("docx_report_module", MODULE_PATH)
docx_report_module = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(docx_report_module)
DocxReportGenerator = docx_report_module.DocxReportGenerator

CLEANER_PATH = Path(__file__).resolve().parents[1] / "tools" / "document" / "file_cleaner.py"
CLEANER_SPEC = importlib.util.spec_from_file_location("file_cleaner_module", CLEANER_PATH)
file_cleaner_module = importlib.util.module_from_spec(CLEANER_SPEC)
assert CLEANER_SPEC and CLEANER_SPEC.loader
CLEANER_SPEC.loader.exec_module(file_cleaner_module)
clean_docx = file_cleaner_module.clean_docx


def _add_inserted_revision(docx_path: Path) -> None:
    """Inject a simple tracked insertion into the first paragraph."""
    with zipfile.ZipFile(docx_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}

    document_xml = files["word/document.xml"]
    doc = Document(docx_path)
    paragraph = doc.paragraphs[0]
    first_run = paragraph.runs[0]._r

    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "76")
    deleted.set(qn("w:author"), "Original Reviewer")
    deleted.set(qn("w:date"), "2026-06-01T08:00:00Z")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "10日内"
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)

    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "77")
    inserted.set(qn("w:author"), "Original Reviewer")
    inserted.set(qn("w:date"), "2026-06-01T09:00:00Z")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "30日内"
    run.append(text)
    inserted.append(run)
    first_run.addnext(deleted)
    first_run.addnext(inserted)

    temp_docx = docx_path.with_suffix(".tmp.docx")
    doc.save(temp_docx)
    with zipfile.ZipFile(temp_docx, "r") as edited:
        files["word/document.xml"] = edited.read("word/document.xml")
    temp_docx.unlink()

    # Keep all other OPC parts exactly as the base file wrote them.
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, blob in files.items():
            target.writestr(name, blob)

    assert document_xml != files["word/document.xml"]


def _add_existing_comment_part(docx_path: Path) -> None:
    """Inject an existing comments part so preservation can be asserted."""
    with zipfile.ZipFile(docx_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}

    rels_name = "word/_rels/document.xml.rels"
    rels = etree.fromstring(files[rels_name])
    rel = etree.SubElement(rels, "Relationship")
    rel.set("Id", "rIdExistingComments")
    rel.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
    )
    rel.set("Target", "comments.xml")
    files[rels_name] = etree.tostring(
        rels,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    content_types_name = "[Content_Types].xml"
    content_types = etree.fromstring(files[content_types_name])
    override = etree.SubElement(content_types, "Override")
    override.set("PartName", "/word/comments.xml")
    override.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    files[content_types_name] = etree.tostring(
        content_types,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    files["word/comments.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:comment w:id="5" w:author="Original Reviewer" w:date="2026-06-01T09:00:00Z">'
        b"<w:p><w:r><w:t>\xe5\x8e\x9f\xe6\x9c\x89\xe6\x89\xb9\xe6\xb3\xa8\xe4\xbf\x9d\xe7\x95\x99</w:t></w:r></w:p>"
        b"</w:comment></w:comments>"
    )

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, blob in files.items():
            target.writestr(name, blob)


def _add_comments_extended_part(docx_path: Path) -> None:
    """Inject a commentsExtended relationship before comments.xml exists."""
    with zipfile.ZipFile(docx_path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}

    rels_name = "word/_rels/document.xml.rels"
    rels = etree.fromstring(files[rels_name])
    rel = etree.SubElement(rels, "Relationship")
    rel.set("Id", "rIdCommentsExtended")
    rel.set(
        "Type",
        "http://schemas.microsoft.com/office/2011/relationships/commentsExtended",
    )
    rel.set("Target", "commentsExtended.xml")
    files[rels_name] = etree.tostring(
        rels,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    content_types_name = "[Content_Types].xml"
    content_types = etree.fromstring(files[content_types_name])
    override = etree.SubElement(content_types, "Override")
    override.set("PartName", "/word/commentsExtended.xml")
    override.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml",
    )
    files[content_types_name] = etree.tostring(
        content_types,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    files["word/commentsExtended.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
        b"</w15:commentsEx>"
    )

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, blob in files.items():
            target.writestr(name, blob)


def _comment_reference_ids(docx_path: Path) -> set[str]:
    with zipfile.ZipFile(docx_path, "r") as package:
        document_xml = etree.fromstring(package.read("word/document.xml"))
    return {
        element.get(qn("w:id"))
        for element in document_xml.findall(f".//{qn('w:commentReference')}")
    }


def _comment_ids(docx_path: Path) -> set[str]:
    with zipfile.ZipFile(docx_path, "r") as package:
        comments_xml = etree.fromstring(package.read("word/comments.xml"))
    return {
        element.get(qn("w:id"))
        for element in comments_xml.findall(f".//{qn('w:comment')}")
    }


class DocxAnnotationPreservationTests(unittest.TestCase):
    def test_annotations_use_original_docx_revision_text_and_beijing_time(self):
        with tempfile.TemporaryDirectory() as tmp:
            contract_path = Path(tmp) / "contract.docx"
            output_path = Path(tmp) / "annotated.docx"

            doc = Document()
            paragraph = doc.add_paragraph()
            paragraph.add_run("甲方应在")
            paragraph.add_run("付款。")
            doc.save(contract_path)
            _add_inserted_revision(contract_path)
            _add_existing_comment_part(contract_path)

            results = [
                {
                    "criterion_id": "C1",
                    "criterion": "付款期限",
                    "issues": [
                        {
                            "issue_id": "I1",
                            "quoted_text": "甲方应在30日内付款。",
                            "comment_text": "付款期限需要核查。",
                        }
                    ],
                }
            ]

            generated = DocxReportGenerator._generate_docx_with_comments(
                str(contract_path),
                results,
                str(output_path),
            )

            self.assertEqual(str(output_path), generated)
            self.assertTrue(os.path.exists(output_path))

            with zipfile.ZipFile(output_path, "r") as output:
                document_xml = output.read("word/document.xml").decode("utf-8")
                comments_xml = output.read("word/comments.xml").decode("utf-8")

            self.assertIn("<w:ins", document_xml)
            self.assertIn("<w:del", document_xml)
            self.assertIn('w:id="77"', document_xml)
            self.assertIn("<w:commentRangeStart", document_xml)
            self.assertIn("<w:highlight", document_xml)
            self.assertIn('w:val="yellow"', document_xml)
            self.assertIn('w:id="5"', comments_xml)
            self.assertIn('w:id="6"', comments_xml)
            self.assertIn("原有批注保留", comments_xml)
            self.assertIn("付款期限需要核查。", comments_xml)
            self.assertIn("+08:00", comments_xml)

    def test_comment_references_target_comments_xml_not_extended_comment_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            contract_path = Path(tmp) / "contract.docx"
            output_path = Path(tmp) / "annotated.docx"

            doc = Document()
            doc.add_paragraph("甲方应付款。")
            doc.save(contract_path)
            _add_comments_extended_part(contract_path)
            _add_existing_comment_part(contract_path)

            results = [
                {
                    "criterion_id": "C1",
                    "criterion": "付款条款",
                    "issues": [
                        {
                            "issue_id": "I1",
                            "quoted_text": "甲方应付款。",
                            "comment_text": "AI新增批注。",
                        }
                    ],
                }
            ]

            DocxReportGenerator._generate_docx_with_comments(
                str(contract_path),
                results,
                str(output_path),
            )

            reference_ids = _comment_reference_ids(output_path)
            comments_ids = _comment_ids(output_path)
            self.assertTrue(reference_ids)
            self.assertTrue(
                reference_ids <= comments_ids,
                f"missing comments.xml ids for references: {reference_ids - comments_ids}",
            )

            with zipfile.ZipFile(output_path, "r") as output:
                comments_xml = output.read("word/comments.xml").decode("utf-8")
            self.assertIn("AI新增批注。", comments_xml)

    def test_multi_paragraph_quoted_text_anchors_to_first_matched_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            contract_path = Path(tmp) / "contract.docx"
            output_path = Path(tmp) / "annotated.docx"

            doc = Document()
            doc.add_paragraph(
                "9.3 于任何情形下，腾讯云不对任何间接的、偶然的、特殊的或惩罚性的损害和赔偿承担责任，"
                "例如利润损失、机会损失、声誉/商誉损失或损害等。"
            )
            doc.add_paragraph(
                "腾讯云在本协议项下所承担的损失赔偿责任不超过损害发生时甲方就该服务在本合同期内"
                "过往所支付的服务费用的总额。"
            )
            doc.save(contract_path)

            results = [
                {
                    "criterion_id": "C9",
                    "criterion": "违约责任",
                    "issues": [
                        {
                            "issue_id": "7.1",
                            "quoted_text": (
                                "9.3 于任何情形下，腾讯云不对任何间接的、偶然的、特殊的或惩罚性的损害和赔偿承担责任，"
                                "例如利润损失、机会损失、声誉/商誉损失或损害等。腾讯云在本协议项下所承担的损失赔偿责任"
                                "不超过损害发生时甲方就该服务在本合同期内过往所支付的服务费用的总额。"
                            ),
                            "comment_text": "责任上限需要复核。",
                        }
                    ],
                }
            ]

            DocxReportGenerator._generate_docx_with_comments(
                str(contract_path),
                results,
                str(output_path),
            )

            events_path = output_path.with_name(f"{output_path.stem}_annotation_events.json")
            events = json.loads(events_path.read_text(encoding="utf-8"))
            self.assertEqual(1, events["annotation_stats"]["exact_matched"])
            self.assertEqual(0, events["annotation_stats"]["unmatched"])
            self.assertEqual("anchored", events["events"][0]["status"])
            self.assertEqual("p1", events["events"][0]["paragraph_path"])

            with zipfile.ZipFile(output_path, "r") as output:
                comments_xml = output.read("word/comments.xml").decode("utf-8")
            self.assertIn("责任上限需要复核。", comments_xml)

    def test_anchored_issue_comment_includes_risk_level_author_spacing_and_highlight(self):
        with tempfile.TemporaryDirectory() as tmp:
            contract_path = Path(tmp) / "contract.docx"
            output_path = Path(tmp) / "annotated.docx"

            doc = Document()
            doc.add_paragraph("乙方应按约履行服务义务。")
            doc.save(contract_path)

            results = [
                {
                    "criterion_id": "C1",
                    "criterion": "服务义务",
                    "issues": [
                        {
                            "issue_id": "1.1",
                            "risk_level": "high",
                            "quoted_text": "乙方应按约履行服务义务。",
                            "comment_text": "建议明确未按约履行时的违约责任。",
                        }
                    ],
                }
            ]

            DocxReportGenerator._generate_docx_with_comments(
                str(contract_path),
                results,
                str(output_path),
            )

            with zipfile.ZipFile(output_path, "r") as output:
                document_xml = output.read("word/document.xml").decode("utf-8")
                comments_xml = output.read("word/comments.xml").decode("utf-8")

            self.assertIn('w:author="AI 条款审查"', comments_xml)
            self.assertIn("风险等级：高", comments_xml)
            self.assertIn("建议明确未按约履行时的违约责任。", comments_xml)
            self.assertIn("<w:highlight", document_xml)
            self.assertIn('w:val="yellow"', document_xml)

    def test_unmatched_issue_does_not_create_unhighlighted_word_comment(self):
        with tempfile.TemporaryDirectory() as tmp:
            contract_path = Path(tmp) / "contract.docx"
            output_path = Path(tmp) / "annotated.docx"

            doc = Document()
            doc.add_paragraph("甲方应按期付款。")
            doc.save(contract_path)

            results = [
                {
                    "criterion_id": "C1",
                    "criterion": "付款条款",
                    "issues": [
                        {
                            "issue_id": "1.1",
                            "risk_level": "medium",
                            "quoted_text": "合同中不存在的引用文本。",
                            "comment_text": "建议明确付款违约责任。",
                        }
                    ],
                }
            ]

            DocxReportGenerator._generate_docx_with_comments(
                str(contract_path),
                results,
                str(output_path),
            )

            with zipfile.ZipFile(output_path, "r") as output:
                self.assertNotIn("word/comments.xml", output.namelist())
                document_xml = output.read("word/document.xml").decode("utf-8")
            self.assertNotIn("<w:commentRangeStart", document_xml)
            self.assertNotIn("<w:highlight", document_xml)

            events_path = output_path.with_name(f"{output_path.stem}_annotation_events.json")
            events = json.loads(events_path.read_text(encoding="utf-8"))
            self.assertEqual(1, events["annotation_stats"]["unmatched"])
            self.assertEqual(0, events["annotation_stats"]["exact_matched"])
            self.assertEqual("medium", events["events"][0]["risk_level"])

    def test_clean_docx_accepts_revisions_and_removes_comments_for_review_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            contract_path = Path(tmp) / "contract.docx"

            doc = Document()
            paragraph = doc.add_paragraph()
            paragraph.add_run("甲方应在")
            paragraph.add_run("付款。")
            doc.save(contract_path)
            _add_inserted_revision(contract_path)
            _add_existing_comment_part(contract_path)

            cleaned_path = Path(clean_docx(contract_path))
            try:
                cleaned_doc = Document(cleaned_path)
                self.assertEqual("甲方应在30日内付款。", cleaned_doc.paragraphs[0].text)
                with zipfile.ZipFile(cleaned_path, "r") as cleaned_zip:
                    self.assertNotIn("word/comments.xml", cleaned_zip.namelist())
            finally:
                if cleaned_path.exists():
                    cleaned_path.unlink()


if __name__ == "__main__":
    unittest.main()
